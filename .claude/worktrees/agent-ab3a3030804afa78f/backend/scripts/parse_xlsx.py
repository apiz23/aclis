# backend/scripts/parse_xlsx.py
"""
Sheet identification and row parsing for the main leaders xlsx.

Before running the full import, discover sheet names:
    py -3.12 -m scripts.run_import --discover --leaders "ID KKG UBAH BETUL .xlsx"
Then update PONTIAN_MUKIM_SHEETS to match the printed names.
"""
from __future__ import annotations
import re
from dateutil import parser as dateutil_parser
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from scripts.models import LeaderRow, RejectRow

# Update after running --discover on the real xlsx
PONTIAN_MUKIM_SHEETS: set[str] = {
    "BENUT", "SERKAT", "PONTIAN", "AYER BALOI", "PEKAN NENAS",
    "SRI GADING", "KUKUP", "RIMBA TERJUN",
}

# Maps lowercased header fragments → LeaderRow field names
_COL_MAP: dict[str, str] = {
    "nama":              "name",
    "no. ic":            "ic_no",
    "no ic":             "ic_no",
    "no.ic":             "ic_no",
    "ic":                "ic_no",
    "kampung":           "kampung_name",
    "kg.":               "kampung_name",
    "tarikh lantikan":   "tarikh_lantikan",
    "tarikh dilantik":   "tarikh_lantikan",
    "parti lantikan":    "parti_lantikan",
    "parti terkini":     "parti_terkini",
    "parti semasa":      "parti_terkini",
}


def list_sheets(xlsx_path: str) -> list[str]:
    """Return sheet names. Run with --discover to check the real xlsx."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def _map_headers(row: tuple) -> dict[str, int]:
    """Map header cell text → column index using _COL_MAP (first-match wins)."""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        if cell is None:
            continue
        text = str(cell).strip().lower()
        for key, field in _COL_MAP.items():
            if key in text and field not in mapping:
                mapping[field] = idx
                break
    return mapping


def _normalize_ic(raw: str) -> str | None:
    """Strip hyphens/spaces; return 12-digit string or None if invalid."""
    cleaned = re.sub(r"[\-\s]", "", str(raw).strip())
    return cleaned if re.match(r"^\d{12}$", cleaned) else None


def _normalize_date(raw) -> str | None:
    """Return ISO 'YYYY-MM-DD' from any common format, or None."""
    if raw is None or str(raw).strip() == "":
        return None
    # openpyxl may return datetime.datetime directly
    if hasattr(raw, "strftime"):
        return raw.strftime("%Y-%m-%d")
    try:
        return dateutil_parser.parse(str(raw), dayfirst=True).strftime("%Y-%m-%d")
    except (ValueError, OverflowError):
        return None


def _parse_sheet(
    ws: Worksheet,
    mukim_name: str,
    source_file: str,
) -> tuple[list[LeaderRow], list[RejectRow]]:
    rows_iter = ws.iter_rows(values_only=True)

    # Find header row: first row whose cells contain 'nama' or 'ic'
    header_map: dict[str, int] = {}
    header_row_num = 0
    for row_num, row in enumerate(rows_iter, start=1):
        m = _map_headers(row)
        if "name" in m or "ic_no" in m:
            header_map = m
            header_row_num = row_num
            break

    if not header_map:
        return [], [RejectRow(source_file, mukim_name, 0, "no recognizable header", "")]

    leaders: list[LeaderRow] = []
    rejects: list[RejectRow] = []

    for row_num, row in enumerate(rows_iter, start=header_row_num + 1):
        if all(v is None or str(v).strip() == "" for v in row):
            continue  # blank row

        def get(field: str):
            idx = header_map.get(field)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        name = str(get("name") or "").strip()
        ic_raw = str(get("ic_no") or "").strip()
        kampung_name = str(get("kampung_name") or "").strip()

        if not name and not ic_raw:
            continue  # silently skip empty data rows

        ic_norm = _normalize_ic(ic_raw)
        if ic_norm is None:
            rejects.append(RejectRow(
                source_file, mukim_name, row_num,
                f"invalid IC: {ic_raw!r}",
                f"name={name}, kampung={kampung_name}",
            ))
            continue

        leaders.append(LeaderRow(
            name=name,
            ic_no=ic_norm,
            type="ketua_kampung",
            kampung_name=kampung_name,
            mukim_name=mukim_name,
            tarikh_lantikan=_normalize_date(get("tarikh_lantikan")),
            parti_lantikan=str(get("parti_lantikan") or "").strip(),
            parti_terkini=str(get("parti_terkini") or "").strip(),
            source_row=row_num,
        ))

    return leaders, rejects


def parse_leaders_file(
    xlsx_path: str,
    include_sheets: set[str] | None = None,
) -> tuple[list[LeaderRow], list[RejectRow]]:
    """Parse all whitelisted Pontian sheets. Returns (leaders, rejects)."""
    whitelist = {s.upper() for s in (include_sheets or PONTIAN_MUKIM_SHEETS)}
    wb = openpyxl.load_workbook(xlsx_path, read_only=False, data_only=True)

    all_leaders: list[LeaderRow] = []
    all_rejects: list[RejectRow] = []

    for sheet_name in wb.sheetnames:
        if sheet_name.strip().upper() not in whitelist:
            continue
        ws = wb[sheet_name]
        leaders, rejects = _parse_sheet(ws, sheet_name.strip().upper(), xlsx_path)
        all_leaders.extend(leaders)
        all_rejects.extend(rejects)

    wb.close()
    return all_leaders, all_rejects
